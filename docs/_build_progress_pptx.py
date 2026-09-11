"""Generate AgriConnect Week 1–3 faculty progress decks (UCS503P)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK = RGBColor(0x14, 0x1A, 0x12)
MUTED = RGBColor(0x3F, 0x4A, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF3, 0xF7, 0xF1)


def set_run(run, size=18, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_bg(slide, color):
    fill = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.333),
        Inches(7.5),
    )
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    fill.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=32):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=DARK)


def add_subtitle(slide, text, top=0.95):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=16, color=MUTED)


def add_bullets(slide, items, top=1.5, left=0.6, width=12.1, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=DARK)


def add_footer(slide, week_label):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"AgriConnect  ·  UCS503P  ·  {week_label}  ·  Kabir Manchanda & Manbhav Kumar Terry"
    set_run(run, size=11, color=MUTED)


def new_slide(prs, week_label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, LIGHT)
    add_bar(slide)
    add_footer(slide, week_label)
    return slide


def title_slide(prs, week_num, dates, milestone, status):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # dark hero
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.5), Inches(1))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "AgriConnect"
    set_run(run, size=44, bold=True, color=WHITE)

    box2 = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11.5), Inches(0.6))
    p = box2.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"Week {week_num} Progress Report"
    set_run(run, size=28, bold=True, color=GREEN)

    box3 = slide.shapes.add_textbox(Inches(0.9), Inches(3.5), Inches(11.5), Inches(1.8))
    tf = box3.text_frame
    lines = [
        "Smart Farm-to-Market & Crop Advisory Platform",
        f"Period: {dates}",
        f"Milestone: {milestone}",
        f"Status: {status}",
        "Team: Kabir Manchanda (1024030415) · Manbhav Kumar Terry (1024030427)",
        "Course: UCS503P · TIET",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = line
        set_run(run, size=16, color=RGBColor(0xC8, 0xD4, 0xC4))


def build_week(path: Path, week_num: int, dates: str, milestone: str, status: str, sections: list):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    label = f"Week {week_num}"
    title_slide(prs, week_num, dates, milestone, status)
    for title, subtitle, bullets in sections:
        slide = new_slide(prs, label)
        add_title(slide, title)
        if subtitle:
            add_subtitle(slide, subtitle)
            add_bullets(slide, bullets, top=1.45)
        else:
            add_bullets(slide, bullets, top=1.2)
    prs.save(path)
    print(f"Wrote {path.name}")


def main():
    build_week(
        OUT / "AgriConnect_Progress_Week1.pptx",
        1,
        "Mon 24 Aug → Mon 31 Aug 2026",
        "Lab — functional CRUD marketplace",
        "Green — lab acceptance criteria met",
        [
            (
                "Agenda",
                None,
                [
                    "Problem & one-line pitch",
                    "Lab milestone goals",
                    "What we delivered this week",
                    "Stack & architecture",
                    "Team split",
                    "Demo checklist",
                    "Next week (Prototype)",
                ],
            ),
            (
                "Problem & pitch",
                "Why AgriConnect exists",
                [
                    "Farmers lack a direct channel to buyers; buyers lack reliable sourcing from growers",
                    "Value leaks to opaque intermediary chains",
                    "Pitch: Direct market access and price transparency for farmers",
                    "Lab focus: production-shaped CRUD marketplace (not a throwaway demo)",
                ],
            ),
            (
                "Lab goals (Week 1)",
                "Faculty-facing Lab milestone",
                [
                    "Auth + RBAC: Farmer / Buyer / Admin",
                    "Produce listings with photo uploads",
                    "Orders with inventory-safe quantity + state machine",
                    "Admin moderation, notifications, basic reports",
                    "Frontend talks only to Express — never to the database",
                ],
            ),
            (
                "Delivered — backend",
                "Kabir — Express / Prisma / Supabase",
                [
                    "Monorepo + Express REST API under /api/v1",
                    "Prisma schema + migrations on Supabase PostgreSQL",
                    "JWT access + HttpOnly refresh cookies; Zod validation",
                    "Listings CRUD + Supabase Storage photos (1–5)",
                    "Order state machine: pending → accepted → confirmed → fulfilled / cancelled",
                    "Admin verify / suspend / moderate; notifications; reports",
                    "Password reset path; rate limits; ownership checks",
                ],
            ),
            (
                "Delivered — frontend",
                "Manbhav — Next.js App Router",
                [
                    "Register / login / session refresh wiring",
                    "Marketplace browse + listing detail",
                    "Farmer listing create / edit / publish flow",
                    "Buyer order place + status views",
                    "Notifications UI + early landing / public profile",
                    "In-memory access token (no localStorage secrets)",
                ],
            ),
            (
                "Stack",
                "Locked Lab stack",
                [
                    "Frontend: Next.js (App Router) + TypeScript + Tailwind",
                    "Backend: Node.js + Express 5 + TypeScript (strict)",
                    "DB: PostgreSQL on Supabase via Prisma",
                    "Storage: Supabase Storage (service role on server only)",
                    "Auth: JWT access + refresh cookie",
                    "Validation: Zod on env + every request",
                ],
            ),
            (
                "Demo checklist",
                "What faculty can see end-to-end",
                [
                    "Register farmer + buyer; login as seeded admin",
                    "Create draft listing → upload photos → publish active",
                    "Buyer places order; farmer accepts → confirms → fulfills",
                    "Admin suspends user / removes listing",
                    "Notifications appear on order events",
                    "Health: GET /health and /ready",
                ],
            ),
            (
                "Next week — Prototype",
                "Deferred from Lab",
                [
                    "Real-time order chat (Socket.io)",
                    "Live mandi prices (Agmarknet / data.gov.in)",
                    "Logistics tracking + escrow-style payments",
                    "Buyer alerts + listing lifecycle jobs",
                    "Docker Compose + GitHub Actions CI",
                    "i18n (Hindi / Punjabi) and guest browse polish",
                ],
            ),
        ],
    )

    build_week(
        OUT / "AgriConnect_Progress_Week2.pptx",
        2,
        "Mon 31 Aug → Mon 7 Sep 2026",
        "Prototype — real-time, production-shaped system",
        "Green — Prototype demo-ready on backend_frontend",
        [
            (
                "Agenda",
                None,
                [
                    "Week 1 recap → Week 2 goals",
                    "Market intelligence (mandi)",
                    "Realtime chat, reviews, assistant",
                    "Logistics, payments stub, alerts",
                    "Docker + CI + i18n",
                    "Verification & carry-over to Week 3",
                ],
            ),
            (
                "From Lab to Prototype",
                "What changed this week",
                [
                    "Lab CRUD marketplace was complete (Week 1)",
                    "Prototype adds realtime, market data, packaging, and trust features",
                    "Goal: faculty-demoable production-shaped system",
                    "Status: Green — core Prototype slice committed (49a11f1 + CI fixes)",
                ],
            ),
            (
                "Market intelligence",
                "Agmarknet / data.gov.in",
                [
                    "Mandi prices service with latestOnly + ~90-day history",
                    "APIs: GET /market/prices, /prices/summary, compare",
                    "Frontend /market-prices page with resilient load UX",
                    "Haryana Potato / Onion demo rows when feed has arrivals",
                    "Listing-vs-mandi compare badge on marketplace",
                ],
            ),
            (
                "Realtime & trust",
                "Chat, reviews, Kisan assistant",
                [
                    "Order-scoped chat (REST + Socket.io) + typing indicators",
                    "Reviews after fulfilled orders",
                    "Kisan AI assistant backend + UI wiring (Gemini when keyed)",
                    "Order timeline component for status history",
                    "Guest browse + landing polish",
                ],
            ),
            (
                "Ops & extras",
                "Logistics, payments stub, alerts, jobs",
                [
                    "PATCH /orders/:id/logistics checkpoints",
                    "Escrow-style payment stub (hold / confirm mock)",
                    "Buyer produce alerts (/buyer/alerts)",
                    "Listing expiry warnings + auto-expire jobs",
                    "V2 schema migration + demo seed data",
                ],
            ),
            (
                "Docker, CI, i18n",
                "Production-shaped packaging",
                [
                    "docker-compose.yml + Dockerfiles for postgres / API / web",
                    "GitHub Actions CI: typecheck, lint, tests",
                    "Hindi / Punjabi (en/hi/pa) language switcher",
                    "Register required fields: phone, state, district",
                    "Localize screens end-to-end; locale cookie → <html lang>",
                ],
            ),
            (
                "Verification",
                "Evidence before close",
                [
                    "Backend Vitest green (43/43 at Week 2 close)",
                    "Frontend build + i18n parity checks",
                    "API_ENDPOINTS.md catalog for boss review",
                    "Remote: origin/backend_frontend",
                ],
            ),
            (
                "Carry into Week 3",
                "Still uncommitted / finishing",
                [
                    "In-app voice calls (WebRTC)",
                    "Payment methods + real Razorpay escrow / webhook",
                    "Anti-bypass contact guard beyond chat",
                    "Expand to 13 Indian locales + frontend test layer",
                    "Faculty Prototype demo from master",
                ],
            ),
        ],
    )

    build_week(
        OUT / "AgriConnect_Progress_Week3.pptx",
        3,
        "Mon 7 Sep → Mon 14 Sep 2026",
        "Prototype close-out + Capstone start (Grounded Kisan)",
        "Green — Capstone Grounded Kisan shipping to master",
        [
            (
                "Agenda",
                None,
                [
                    "Week overview & phase snapshot",
                    "Prototype hardening shipped",
                    "Capstone: Grounded Kisan",
                    "i18n + tests + proposal",
                    "Risks & backlog",
                    "Next week / faculty demo",
                ],
            ),
            (
                "Phase snapshot",
                "Lab → Prototype → Capstone",
                [
                    "Lab: Complete (Week 1) — CRUD marketplace",
                    "Prototype: Complete (Week 2–3) — realtime, payments, calls, 13 locales",
                    "Capstone: In progress — Grounded Kisan RAG + agronomist + weather",
                    "Focus: faculty demo from master; then Capstone polish",
                ],
            ),
            (
                "Prototype hardening",
                "Shipped this week",
                [
                    "Payments: UPI / card / netbanking / COD; Razorpay verify + signed webhook",
                    "Contact guard on chat, order notes, listing description/variety/village",
                    "In-app voice calls (WebRTC + Socket.io); 45s ring timeout; smoke 13/13",
                    "Edge hardening: 413 mapping, webhook checkout match, socket token refresh",
                    "Natural orders list: active statuses above cancelled",
                    "Pushed backend_frontend → merged master (local: 138 backend / 128 frontend tests)",
                ],
            ),
            (
                "Capstone — Grounded Kisan",
                "PRD §8.2 advisory platform start",
                [
                    "Same Kisan widget; curated knowledge pack + cite / refuse",
                    "AGRONOMIST role + AdvisoryEscalation queue for high-stakes asks",
                    "Profile weather + 5-day forecast (OpenWeather)",
                    "Faster talk: pinned TTS model; short speak clips",
                    "Offline eval: eval:grounded 30/30 on curated pack",
                ],
            ),
            (
                "Product & docs",
                "Also this week",
                [
                    "13-locale UI: 477 keys × 13; Vitest parity / width / no-hardcoded-strings",
                    "Header RTL + Noto webfonts for Indic scripts / Urdu",
                    "project-proposal/ on master (PDF + LaTeX; submitted to Mr. Hardik)",
                    "UCS503 faculty docs skeleton: index, requirements, architecture, testing, evaluation, ROADMAP",
                ],
            ),
            (
                "Risks & backlog",
                "Honest remaining work",
                [
                    "Native-speaker pass on machine-translated locales",
                    "TURN server for calls behind strict NATs (STUN only today)",
                    "Farmer payouts / settlement (escrow hold/refund done; payout not)",
                    "Contact blocking on review comments (last free-text gap)",
                    "Razorpay webhook tunnel needed for full local webhook demos",
                    "Capstone backlog: ML price prediction, crop disease CV",
                ],
            ),
            (
                "Faculty demo path",
                "Suggested walkthrough",
                [
                    "1. Landing → marketplace browse → mandi prices (Haryana Potato/Onion/Tomato)",
                    "2. Place order → chat → voice call (no phone exchange)",
                    "3. Payment method + escrow confirm path",
                    "4. Farmer Kisan: grounded answer with citations; refuse / escalate",
                    "5. Switch language; show admin moderate briefly",
                ],
            ),
            (
                "Next week",
                "Week 4 outlook",
                [
                    "Faculty Prototype + Grounded Kisan demo from master",
                    "Optional: TURN for calls; farmer payout settlement",
                    "Optional: expand grounded-eval; native locale pass",
                    "Capstone: ML price / crop CV remain backlog unless prioritized",
                ],
            ),
        ],
    )


if __name__ == "__main__":
    main()
